"""Импорт предписаний в реестр из таблицы CSV/XLSX.

Загрузка возможна инспектором/администратором (веб-маршрут
/prescriptions/import). Формат — строка заголовков + данные.

Заголовки принимаются в двух вариантах: канонические англо-имена
(doc_number, doc_date, ...) и русская раскладка окна «Реестр предписаний»
(№, Дата выдачи, Дата устранения, Тип, Объект, Подрядчик, Категория, Вид,
Описание нарушений, ФИО отв. за устр., Инспектор, Статус, Срок).
Регистр и лишние пробелы в заголовках не важны; пустые колонки-хвосты
игнорируются, полностью пустые строки пропускаются.

Обязательные колонки: doc_number (№), doc_date (Дата выдачи), type (Тип).
Пустые object/contractor заполняются справочником «Не указан».
Пустое описание нарушений берется из обоснования, иначе «Без описания».

type: internal/customer (допустимы «внутреннее»/«заказчика»);
status: значения enum или русские подписи («Закрыто», «В работе»...);
даты: ISO, ДД.ММ.ГГГГ, datetime из Excel; в текстовых сроках («до 30.10.2021»)
дата извлекается по шаблону, нераспознанное — пустое поле + warning.

Поведение импорта (стиль users.service.import_users_from_csv):
  * справочники (object/contractor/enterprise/category/kind/mgmt/issuer)
    ищутся по имени, отсутствующие создаются;
  * status из файла — режим переноса истории: назначается напрямую с
    записью в историю (imported_status); для closed проставляется closed_at
    (из файла или текущее время);
  * дубли (doc_number + doc_date + object + type) пропускаются — импорт
    идемпотентен;
  * ошибка в строке не прерывает импорт: строка откатывается, ошибка
    собирается в summary.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from src.db.models import (
    Contractor,
    ConstructionManagement,
    Enterprise,
    Object,
    Prescription,
    PrescriptionCategory,
    PrescriptionInspector,
    PrescriptionKind,
    PrescriptionStatus,
    PrescriptionType,
    User,
    UserRole,
)
from src.db.repositories import find_prescription_duplicate
from src.prescriptions import service as prescriptions_service
from src.prescriptions.schemas import PrescriptionCreate

IMPORT_COLUMNS = [
    "doc_number", "doc_date", "type", "object", "contractor",
    "enterprise", "category", "kind", "mgmt", "issuer",
    "description", "justification",
    "deadline", "issue_date", "actual_fix_date",
    "status", "closed_at", "responsible", "inspector",
]
REQUIRED_COLUMNS = {"doc_number", "doc_date", "type"}

UNSPECIFIED = "Не указан"
NO_DESCRIPTION = "Без описания"

# Русская раскладка окна «Реестр предписаний» -> канонические имена колонок.
HEADER_ALIASES = {
    "№": "doc_number",
    "номер": "doc_number",
    "№ документа": "doc_number",
    "документ": "doc_number",
    "дата выдачи": "doc_date",
    "дата составления": "issue_date",
    "дата устранения": "actual_fix_date",
    "тип": "type",
    "объект": "object",
    "подрядчик": "contractor",
    "категория": "category",
    "вид": "kind",
    "описание нарушений": "description",
    "описание": "description",
    "фио отв. за устр.": "responsible",
    "фио отв. за устранение": "responsible",
    "ответственный": "responsible",
    "инспектор": "inspector",
    "статус": "status",
    "срок": "deadline",
    "предприятие": "enterprise",
    "обоснование": "justification",
    "управление": "mgmt",
    "выдавший": "issuer",
    "дата закрытия": "closed_at",
}

# Русские подписи статусов (обратная карта STATUS_LABELS из веб-слоя).
_STATUS_ALIASES = {
    "черновик (на подпись)": PrescriptionStatus.DRAFT_FOR_SIGNATURE,
    "черновик": PrescriptionStatus.DRAFT_FOR_SIGNATURE,
    "новое": PrescriptionStatus.NEW,
    "на проверке": PrescriptionStatus.UNDER_REVIEW,
    "в работе": PrescriptionStatus.IN_WORK,
    "заявлено об устранении": PrescriptionStatus.FIXED_PENDING_REVIEW,
    "закрыто": PrescriptionStatus.CLOSED,
    "отклонено": PrescriptionStatus.REJECTED,
}

_TYPE_ALIASES = {
    "internal": "internal",
    "внутреннее": "internal",
    "внутренний": "internal",
    "customer": "customer",
    "заказчика": "customer",
    "заказчик": "customer",
}

_NAME_REF_MODELS = {
    "enterprise": Enterprise,
    "category": PrescriptionCategory,
    "kind": PrescriptionKind,
    "mgmt": ConstructionManagement,
    "issuer": PrescriptionInspector,
}

_DATE_IN_TEXT = re.compile(r"(\d{1,2})[./](\d{1,2})[./](\d{4})")
_ISO_IN_TEXT = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})")


# --- заголовки ---------------------------------------------------------------------


def _canon_header(raw: Any) -> str:
    key = re.sub(r"\s+", " ", str(raw or "").strip().lower())
    if key in IMPORT_COLUMNS:
        return key
    return HEADER_ALIASES.get(key, key)


# --- чтение файла ----------------------------------------------------------------


def read_table(data: bytes, filename: str) -> list[dict]:
    """Разбор загруженного файла в список строк-словарей (ключи — канонические)."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _read_csv(data)
    if name.endswith(".xlsx"):
        return _read_xlsx(data)
    raise ValueError("Поддерживаются файлы .csv и .xlsx")


def _read_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if reader.fieldnames is None:
        raise ValueError("Пустой файл")
    columns = [_canon_header(h) for h in reader.fieldnames]
    _require_columns(columns)
    rows = []
    for row in reader:
        clean = {}
        for column, value in zip(columns, (row.get(orig) for orig in reader.fieldnames)):
            if not column:
                continue
            clean[column] = value if value is not None else ""
        if any(str(v).strip() for v in clean.values()):
            rows.append(clean)
    return rows


def _read_xlsx(data: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("Для импорта XLSX нужен пакет openpyxl") from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    it = sheet.iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        raise ValueError("Пустой лист")
    columns = [_canon_header(c) for c in header]
    _require_columns(columns)
    rows = []
    for values in it:
        values = list(values or [])
        row = {}
        for idx, column in enumerate(columns):
            if not column:
                continue
            value = values[idx] if idx < len(values) else None
            row[column] = value
        if any(_norm_cell(v) != "" for v in row.values()):
            rows.append(row)
    wb.close()
    return rows


def _require_columns(columns: list[str]) -> None:
    if not REQUIRED_COLUMNS <= set(columns):
        raise ValueError(
            "Обязательные колонки: № (doc_number), Дата выдачи (doc_date), Тип (type); найдено: "
            + ", ".join(c for c in dict.fromkeys(columns) if c)
        )


def _norm_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# --- разбор значений ---------------------------------------------------------------


def _cell(row: dict, key: str):
    value = row.get(key)
    if isinstance(value, str):
        return value.strip()
    return value if value is not None else ""


def _to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    if not raw:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return date.fromisoformat(raw)


def _to_datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    raw = str(value or "").strip()
    if not raw:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%d.%m.%Y"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    raise ValueError(f"Не распознано время: {raw!r}")


def _to_date_forgiving(value, field: str, doc_number: str, warnings: list[str]) -> date | None:
    """Мягкий разбор даты: пустое -> None; текст с датой («до 30.10.2021») -> дата;
    не распознано -> None + warning (строка импортируется)."""
    if isinstance(value, (datetime, date)):
        return _to_date(value)
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        return _to_date(raw)
    except ValueError:
        pass
    match = _DATE_IN_TEXT.search(raw) or _ISO_IN_TEXT.search(raw)
    if match:
        try:
            if len(match.group(1)) == 4:
                return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
            return date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
        except ValueError:
            pass
    warnings.append(f"{doc_number}: {field} «{raw}» не распознано — оставлено пустым")
    return None


def _parse_type(raw: str) -> PrescriptionType:
    key = (raw or "").strip().lower()
    if key not in _TYPE_ALIASES:
        raise ValueError(f"Неизвестный тип {raw!r}; ожидается internal/customer")
    return PrescriptionType(_TYPE_ALIASES[key])


def _parse_status(raw: str) -> PrescriptionStatus | None:
    key = re.sub(r"\s+", " ", (raw or "").strip().lower())
    if not key:
        return None
    try:
        return PrescriptionStatus(key)
    except ValueError:
        pass
    if key in _STATUS_ALIASES:
        return _STATUS_ALIASES[key]
    raise ValueError(
        f"Неизвестный статус {raw!r}; ожидается одно из: "
        + ", ".join(s.value for s in PrescriptionStatus)
        + " или "
        + ", ".join(sorted(_STATUS_ALIASES))
    )


# --- разрешение справочников ---------------------------------------------------------


def _resolve_or_create(db: Session, model, name: str, new_refs: list[str]) -> int:
    if not name:
        raise ValueError("Пустое имя справочника")
    existing = db.scalar(select(model).where(model.name == name))
    if existing is not None:
        return existing.id
    created = model(name=name)
    db.add(created)
    db.flush()
    new_refs.append(f"{model.__name__}: {name}")
    return created.id


def _find_user_by_name(db: Session, full_name: str, role: UserRole) -> int | None:
    return db.scalar(
        select(User.id)
        .where(User.full_name == full_name, User.role == role)
        .order_by(User.id)
    )


# --- основной импорт ------------------------------------------------------------------


def import_prescriptions(db: Session, rows: list[dict], *, actor_user_id: int | None) -> dict:
    """Импортировать строки таблицы в реестр. Возвращает сводку."""
    if actor_user_id is None:
        raise PermissionError("Импорт доступен только авторизованному инспектору/администратору")
    actor = db.get(User, actor_user_id)
    if actor is None or not actor.is_active:
        raise PermissionError("Пользователь не найден или неактивен")
    if actor.role not in (UserRole.INSPECTOR, UserRole.ADMIN):
        raise PermissionError("Импорт в реестр доступен инспектору или администратору")

    summary: dict = {
        "created": [], "skipped": [], "errors": [], "new_refs": [],
        "no_responsible": [], "no_inspector": [], "warnings": [],
    }

    for line_no, raw_row in enumerate(rows, start=2):
        row = {key: _norm_cell(value) for key, value in raw_row.items()}
        doc_number = str(_cell(row, "doc_number") or "")
        try:
            if not doc_number:
                raise ValueError("Пустая обязательная колонка № (doc_number)")
            if not str(_cell(row, "doc_date") or ""):
                raise ValueError("Пустая обязательная колонка Дата выдачи (doc_date)")

            ptype = _parse_type(_cell(row, "type"))
            doc_date = _to_date(_cell(row, "doc_date"))
            if doc_date is None:
                raise ValueError("Не распознана Дата выдачи (doc_date)")

            object_name = str(_cell(row, "object") or "") or UNSPECIFIED
            contractor_name = str(_cell(row, "contractor") or "") or UNSPECIFIED
            object_id = _resolve_or_create(db, Object, object_name, summary["new_refs"])
            contractor_id = _resolve_or_create(db, Contractor, contractor_name, summary["new_refs"])

            ref_ids: dict[str, int | None] = {}
            for column, model in _NAME_REF_MODELS.items():
                name = str(_cell(row, column) or "")
                ref_ids[column] = (
                    _resolve_or_create(db, model, name, summary["new_refs"]) if name else None
                )

            responsible_name = str(_cell(row, "responsible") or "")
            responsible_id = (
                _find_user_by_name(db, responsible_name, UserRole.CONTRACTOR_RESPONSIBLE)
                if responsible_name else None
            )
            if responsible_name and responsible_id is None:
                summary["no_responsible"].append(f"{doc_number}: {responsible_name}")

            inspector_name = str(_cell(row, "inspector") or "")
            inspector_id = (
                _find_user_by_name(db, inspector_name, UserRole.INSPECTOR)
                if inspector_name else None
            )
            if inspector_name and inspector_id is None:
                summary["no_inspector"].append(f"{doc_number}: {inspector_name}")

            warnings: list[str] = summary["warnings"]
            deadline = _to_date_forgiving(_cell(row, "deadline"), "Срок", doc_number, warnings)
            issue_date = _to_date_forgiving(_cell(row, "issue_date"), "Дата составления", doc_number, warnings)
            actual_fix = _to_date_forgiving(_cell(row, "actual_fix_date"), "Дата устранения", doc_number, warnings)

            justification = str(_cell(row, "justification") or "") or None
            description = str(_cell(row, "description") or "") or justification or NO_DESCRIPTION

            target_status = _parse_status(_cell(row, "status"))

            data = PrescriptionCreate(
                prescription_type=ptype,
                doc_number=doc_number,
                doc_date=doc_date,
                issue_date=issue_date,
                actual_fix_date=actual_fix,
                enterprise_id=ref_ids["enterprise"],
                object_id=object_id,
                contractor_id=contractor_id,
                description=description,
                justification=justification,
                deadline=deadline,
                category_id=ref_ids["category"],
                kind_id=ref_ids["kind"],
                mgmt_id=ref_ids["mgmt"],
                issuer_id=ref_ids["issuer"],
                responsible_user_id=responsible_id,
                inspector_user_id=inspector_id,
            )

            duplicate = find_prescription_duplicate(
                db,
                doc_number=data.doc_number,
                doc_date=data.doc_date,
                object_id=data.object_id,
                prescription_type=data.prescription_type,
            )
            if duplicate is not None:
                summary["skipped"].append(doc_number)
                continue

            prescription = prescriptions_service.create_prescription(db, data, actor_user_id=actor.id)

            if target_status is not None and target_status != prescription.status:
                prescription.status = target_status
                if target_status == PrescriptionStatus.CLOSED:
                    closed_at = _to_datetime(_cell(row, "closed_at")) or prescriptions_service._now()
                    prescription.closed_at = closed_at
                    prescription.closed_by_user_id = actor.id
                    if prescription.confirmed_at is None:
                        prescription.confirmed_at = closed_at
                    prescription.requires_review = False
                prescriptions_service._add_history(
                    db,
                    prescription_id=prescription.id,
                    actor_id=actor.id,
                    change_type="imported_status",
                    new_value={"status": target_status.value},
                )
                db.commit()

            summary["created"].append(doc_number)
        except Exception as exc:
            db.rollback()
            summary["errors"].append(f"строка {line_no} ({doc_number or 'без №'}): {exc}")

    return summary


def template_csv() -> str:
    """CSV-шаблон для скачивания (русская реестровая раскладка заголовков)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([
        "№", "Дата выдачи", "Дата устранения", "Тип", "Объект", "Подрядчик",
        "Категория", "Вид", "Описание нарушений", "Обоснование",
        "ФИО отв. за устр.", "Инспектор", "Статус", "Срок", "Предприятие",
        "Дата составления",
    ])
    return buffer.getvalue()
